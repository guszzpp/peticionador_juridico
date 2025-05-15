# src/peticionador/agentes/agente_gerador_peca.py

import logging
from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime
from odf.opendocument import OpenDocumentText
from odf.text import P as OdfParagraph
from docx import Document
from peticionador.servicos.integrador_gemini import ClienteGemini
from peticionador.utilitarios.substituidor_docx import substituir_placeholders_em_docx

RAIZ_PROJETO = Path(__file__).resolve().parents[3]
PASTA_SAIDA = RAIZ_PROJETO / "arquivos_gerados"
PASTA_SAIDA.mkdir(parents=True, exist_ok=True)

log = logging.getLogger(__name__)

from docx import Document

def extrair_texto_docx_para_prompt(caminho: Path) -> str:
    """
    Extrai o texto de um arquivo .docx para uso como modelo no prompt da IA.
    Ignora parágrafos vazios e preserva a ordem.
    """
    doc = Document(str(caminho))
    return "\n\n".join(
        par.text.strip()
        for par in doc.paragraphs
        if par.text.strip()
    )

def construir_minuta_com_ia(
    resumo_tecnico: str,
    teses_selecionadas: List[str],
    modelo_base_path: Path,
    dados_processo: Dict[str, str],
    temperatura_ia: float = 0.2,
    max_tokens_ia: Optional[int] = 8000
) -> str:
    log.info(f"Iniciando construção da minuta com IA. Modelo base: {modelo_base_path.name}")
    log.info(f"Teses selecionadas: {teses_selecionadas}")
    log.debug(f"Resumo técnico: {resumo_tecnico[:300]}...")

    try:
        with open(modelo_base_path, "r", encoding="utf-8") as f:
            conteudo_modelo_base = extrair_texto_docx_para_prompt(modelo_base_path)
    except Exception as e:
        log.error(f"Erro ao ler modelo base {modelo_base_path}: {e}")
        return f"[ERRO INTERNO: Falha ao ler modelo base - {e}]"

    teses_prompt = "\n".join(f"- {t.strip()}" for t in teses_selecionadas) or "Nenhuma tese específica foi selecionada pelo usuário para esta seção."

    prompt = f"""
Você é um Procurador de Justiça do MPGO. Construa uma peça completa de contrarrazões conforme o modelo abaixo:

--- MODELO BASE ---
{conteudo_modelo_base}
--- FIM DO MODELO BASE ---

--- RESUMO TÉCNICO ---
{resumo_tecnico}

--- TESES SELECIONADAS ---
{teses_prompt}

--- INSTRUÇÕES ---
1. Substitua o placeholder {{{{RESUMO_PARA_A_PECA}}}} com um parágrafo conciso relatando os fundamentos do recurso.
2. Substitua {{{{TESES_E_ARGUMENTOS}}}} com seções numeradas, uma para cada tese listada, no seguinte formato:
   "1. Da ausência de prequestionamento", "2. Da ofensa reflexa à Constituição", etc.
   Desenvolva cada uma em parágrafos dissertativos.
   Após a última tese, adicione obrigatoriamente uma seção final intitulada:
   "[n+1]. Da conclusão", que traga o fecho da peça com pedido de não provimento do recurso.
3. Substitua os demais placeholders com os dados abaixo ou deixe o placeholder intacto se faltarem dados.
4. Não utilize "Resp" ou "RE", escreva por extenso "Recurso Especial" ou "Recurso Extraordinário".
--- DADOS DO PROCESSO ---
{dados_processo}

Gere apenas o texto final da peça. Nada mais.
    """

    cliente_ia = ClienteGemini(model_name=ClienteGemini.DEFAULT_MODEL_PRO)
    log.info("Enviando prompt à IA...")
    resposta = cliente_ia.gerar_conteudo(prompt, temperatura=temperatura_ia, max_tokens=max_tokens_ia)

    if not resposta or "[ERRO:" in resposta:
        log.error(f"Erro na resposta da IA: {resposta}")
        return f"[FALHA NA GERAÇÃO PELA IA: {resposta}]"

    resposta = resposta.strip()
    log.info("Minuta textual gerada com sucesso.")

    # Nome base e caminhos
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome_base = f"minuta_gerada_{timestamp}"
    caminho_txt = PASTA_SAIDA / f"{nome_base}.txt"
    caminho_docx_saida = PASTA_SAIDA / f"{nome_base}.docx"
    caminho_odt = PASTA_SAIDA / f"{nome_base}.odt"

    arquivos_gerados_nesta_etapa: Dict[str, str] = {}

    # Salva .txt
    with open(caminho_txt, "w", encoding="utf-8") as f_out:
        f_out.write(resposta)
    arquivos_gerados_nesta_etapa["minuta_gerada"] = str(caminho_txt.relative_to(RAIZ_PROJETO))
    log.info(f"Minuta .txt salva em '{caminho_txt.name}'")

    # .docx (preserva formatação)
    caminho_modelo_docx = modelo_base_path.with_suffix(".docx")
    caminho_docx_saida = PASTA_SAIDA / f"{nome_base}.docx"
    if caminho_modelo_docx.exists():
        try:
            substituir_placeholders_em_docx(
                caminho_modelo=caminho_modelo_docx,
                caminho_saida=caminho_docx_saida,
                substituicoes=dados_processo | {
                    "RESUMO_PARA_A_PECA": extrair_placeholder(resposta, "RESUMO_PARA_A_PECA"),
                    "TESES_E_ARGUMENTOS": extrair_placeholder(resposta, "TESES_E_ARGUMENTOS")
                }
            )
            arquivos_gerados_nesta_etapa["minuta_gerada_docx"] = str(caminho_docx_saida.relative_to(RAIZ_PROJETO))
            log.info(f"Minuta .docx salva em '{caminho_docx_saida.name}'")
        except Exception as e_docx:
            log.error(f"Erro ao gerar .docx: {e_docx}", exc_info=True)
    else:
        log.warning(f"Modelo .docx não encontrado: {caminho_modelo_docx.name}")

    # .odt (simples)
    try:
        doc_odt = OpenDocumentText()
        for par in resposta.strip().split("\n\n"):
            doc_odt.text.addElement(OdfParagraph(text=par.strip()))
        doc_odt.save(str(caminho_odt))
        arquivos_gerados_nesta_etapa["minuta_gerada_odt"] = str(caminho_odt.relative_to(RAIZ_PROJETO))
        log.info(f"Minuta .odt salva em '{caminho_odt.name}'")
    except Exception as e_odt:
        log.error(f"Erro ao gerar .odt: {e_odt}", exc_info=True)

    # Registra arquivos para download
    from flask import current_app
    if current_app and current_app.config.get("ULTIMO_PROCESSAMENTO"):
        current_app.config["ULTIMO_PROCESSAMENTO"].setdefault("arquivos", {}).update(arquivos_gerados_nesta_etapa)

    return resposta


def extrair_placeholder(texto: str, nome: str) -> str:
    """Extrai o conteúdo de um placeholder específico, delimitado por {{NOME}} até o próximo {{ ou fim."""
    import re
    pattern = re.compile(rf"\{{{{{nome}}}}}\s*(.*?)(?=\n\s*\{{{{|\Z)", re.DOTALL)
    match = pattern.search(texto)
    return match.group(1).strip() if match else ""
