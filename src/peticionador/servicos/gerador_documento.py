#  src/peticionador/servicos/gerador_documento.py

import os
from datetime import datetime
from typing import Dict

from docx import Document
from docx.shared import Pt
from odf.opendocument import OpenDocumentText
from odf.text import P, Span


def preencher_placeholders(texto_modelo: str, dados: Dict[str, str]) -> str:
    for chave, valor in dados.items():
        texto_modelo = texto_modelo.replace(f"{{{{{chave}}}}}", valor)
    return texto_modelo


def gerar_docx(texto_formatado: str, caminho_saida: str):
    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)

    for paragrafo in texto_formatado.split("\n\n"):
        doc.add_paragraph(paragrafo.strip())

    doc.save(caminho_saida)


def gerar_odt(texto_formatado: str, caminho_saida: str):
    odt = OpenDocumentText()
    for paragrafo in texto_formatado.split("\n\n"):
        p = P()
        span = Span(text=paragrafo.strip())
        p.addElement(span)
        odt.text.addElement(p)
    odt.save(caminho_saida)


def gerar_peca_personalizada(
    estado_peticao, modelo_path: str, saida_dir: str = "arquivos_gerados"
) -> Dict[str, str]:
    if not os.path.exists(saida_dir):
        os.makedirs(saida_dir)

    with open(modelo_path, "r", encoding="utf-8") as f:
        texto_modelo = f.read()

    dados_substituiveis = {
        "NUM_PROCESSO": "0000000-00.0000.0.00.0000",  #  Placeholder manual
        "NOME_RECORRENTE": estado_peticao.estrutura_base.get(
            "recorrente", "Não informado"
        ),
        "RESUMO_TECNICO": estado_peticao.resumo or "[Resumo técnico não gerado]",
        "TESES_E_ARGUMENTOS": "\n\n".join(
            f"- {arg}" for arg in estado_peticao.argumentos_reutilizaveis
        ),
        "NOME_PROMOTOR": "Fulano de Tal",  #  Poderá ser configurável depois
        "CIDADE": "Goiânia",
        "DATA_ATUAL": datetime.today().strftime("%d de %B de %Y"),
    }

    texto_final = preencher_placeholders(texto_modelo, dados_substituiveis)

    caminho_docx = os.path.join(saida_dir, "contrarrazoes_respe.docx")
    caminho_odt = os.path.join(saida_dir, "contrarrazoes_respe.odt")

    gerar_docx(texto_final, caminho_docx)
    gerar_odt(texto_final, caminho_odt)

    return {"docx": caminho_docx, "odt": caminho_odt}
